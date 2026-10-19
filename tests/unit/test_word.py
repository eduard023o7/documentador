"""Word and PDF conversion: BR-U0-27 to BR-U0-30."""

import io
import shutil
import zipfile

import pytest
from docx import Document as WordDocument

from documentador.core.document import parse_document
from documentador.core.profile import load_profile
from documentador.publish.word import (
    WATERMARK_TEXT,
    ConversionError,
    default_link,
    document_markdown,
    draft_template,
    is_draft,
    link_references,
    render_docx,
    render_pdf,
)

from ..conftest import MARCO, ROOT

TEMPLATE = MARCO / "plantilla-word" / "corporativa.docx"
FIXTURE = ROOT / "src" / "documentador" / "publish" / "ejemplo.md"
needs_pandoc = pytest.mark.skipif(not shutil.which("pandoc"), reason="Pandoc no instalado")
needs_office = pytest.mark.skipif(
    not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice no instalado"
)


@pytest.fixture(scope="module")
def doc():
    profile = load_profile(MARCO).profile
    return profile, parse_document(FIXTURE.read_text(encoding="utf-8"), profile)


def _xml(docx: bytes, prefix: str) -> str:
    with zipfile.ZipFile(io.BytesIO(docx)) as z:
        return "".join(z.read(n).decode("utf-8") for n in z.namelist() if n.startswith(prefix))


# BR-U0-30
def test_references_become_links():
    out = link_references("Ver [[LIN-INF-001]] y [[LIN-INF-001.R02]].", default_link)
    assert out == "Ver [LIN-INF-001](LIN-INF-001.md) y [LIN-INF-001.R02](LIN-INF-001.md#lin-inf-001r02)."


def test_document_markdown(doc):
    profile, parsed = doc
    md = document_markdown(parsed, profile)
    assert md.startswith("# EST-INF-AWS-001 · Landing zone en AWS")
    assert "| Estándar | INF | 0.0.0 | en_validacion |" in md
    assert "[[" not in md and "---\nid:" not in md


def test_is_draft(doc):
    _, parsed = doc
    assert is_draft(parsed)


# BR-U0-28
def test_draft_template_has_watermark(tmp_path):
    target = draft_template(TEMPLATE, tmp_path / "borrador.docx")
    xml = _xml(target.read_bytes(), "word/header")
    assert f'string="{WATERMARK_TEXT}"' in xml
    assert WordDocument(str(target)).sections  # still a valid Word file


@needs_pandoc
def test_docx_uses_template_and_links(doc):
    profile, parsed = doc
    docx = render_docx(document_markdown(parsed, profile), TEMPLATE, draft=False)
    word = WordDocument(io.BytesIO(docx))
    assert word.sections[0].header.paragraphs[0].text.startswith("Plantilla genérica")  # BR-U0-27
    assert any(p.style.name.startswith("Heading") for p in word.paragraphs)
    assert "LIN-INF-001.md" in _xml(docx, "word/_rels/document.xml.rels")  # BR-U0-30
    assert WATERMARK_TEXT not in _xml(docx, "word/header")


@needs_pandoc
def test_draft_docx_has_watermark(doc):
    profile, parsed = doc
    docx = render_docx(document_markdown(parsed, profile), TEMPLATE, draft=True)
    assert WATERMARK_TEXT in _xml(docx, "word/header")


@needs_pandoc
@needs_office
def test_pdf_from_docx(doc):
    profile, parsed = doc
    pdf = render_pdf(render_docx(document_markdown(parsed, profile), TEMPLATE, draft=True))
    assert pdf.startswith(b"%PDF")


def test_missing_binary(monkeypatch):
    monkeypatch.setattr("documentador.publish.word.shutil.which", lambda _: None)
    with pytest.raises(ConversionError, match="pandoc"):
        render_docx("# x", TEMPLATE, draft=False)
    with pytest.raises(ConversionError, match="LibreOffice"):
        render_pdf(b"x")
